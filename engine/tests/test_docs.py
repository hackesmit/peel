"""The documents family: PDF, documents, slides, spreadsheets and subtitles.

Every fixture is built at run time by the library that owns the format
(python-docx, python-pptx, openpyxl, PyMuPDF), and every output is re-opened by
the library that would open it for real, because "the conversion exited 0" says
nothing about whether Word, Excel or PowerPoint can read the file. Where the
source format has no Python writer (PPT, XLS, ODS, ODT, RTF) the fixture is
made by converting one that does, which incidentally proves the legacy input
path works.

Tests needing a backend skip when it is missing, so the suite still passes on a
machine where LibreOffice or pandoc is not installed yet. On hqbox and on a
finished George install, nothing here skips.
"""

from __future__ import annotations

import json
import os
import re
import subprocess
import sys
import zipfile
from pathlib import Path

import pytest

from peel_engine import backends
from peel_engine.families import docs
from peel_engine.registry import CONVERT, TOOL, load_registry

ENGINE_DIR = Path(__file__).resolve().parents[1]

# --- skip helpers -----------------------------------------------------------


def needs(*names: str):
    """Skip unless every backend is installed."""
    missing = backends.missing(tuple(names))
    return pytest.mark.skipif(bool(missing), reason=f"needs {', '.join(missing)}")


needs_pandoc = needs("pandoc")
needs_soffice = needs("soffice")
needs_office = needs("pandoc", "soffice")
needs_pymupdf = needs("pymupdf")
needs_pdf2docx = needs("pdf2docx")


@pytest.fixture(scope="module")
def registry():
    return load_registry(force=True)


# --- fixtures: sources built by the library that owns the format ------------


@pytest.fixture(scope="module")
def samples(tmp_path_factory) -> dict[str, Path]:
    """One source file per format, built once for the whole module."""
    pytest.importorskip("docx")
    pytest.importorskip("pptx")
    pytest.importorskip("openpyxl")
    pytest.importorskip("pymupdf")
    folder = tmp_path_factory.mktemp("samples")
    made = {
        "docx": _make_docx(folder / "report.docx"),
        "picdocx": _make_docx_with_picture(folder / "illustrated.docx"),
        "md": _write(folder / "notes.md", "# Notes\n\nSome *emphasis* here.\n\n- one\n- two\n"),
        "txt": _write(folder / "plain.txt", "First line.\nSecond line.\n"),
        "pptx": _make_pptx(folder / "deck.pptx", slides=2),
        "xlsx": _make_xlsx(folder / "book.xlsx"),
        "csv": _write(folder / "table.csv", "name,qty,code\napple,3,00731\npear,5,00942\n"),
        "tsv": _write(folder / "table.tsv", "name\tqty\napple\t3\npear\t5\n"),
        "pdf": _make_pdf(folder / "paper.pdf"),
        "onepage": _make_pdf(folder / "single.pdf", pages=1),
        "srt": _write(
            folder / "movie.srt",
            "1\n00:00:01,000 --> 00:00:03,500\nHello there.\n\n"
            "2\n00:00:04,000 --> 00:00:06,250\n<i>Second</i> line\nwrapped\n\n"
            "3\n00:00:07,000 --> 00:00:08,000\nSecond line\nwrapped\n",
        ),
        "vtt": _write(
            folder / "movie.vtt",
            "WEBVTT\n\nNOTE a comment that even contains --> in it\nsecond note line\n\n"
            "cue-1\n00:01.000 --> 00:03.500 line:90% align:middle\nHello there.\n\n"
            "00:00:04.000 --> 00:00:06.250\nSecond line\n",
        ),
    }
    return made


@pytest.fixture(scope="module")
def legacy(tmp_path_factory, samples) -> dict[str, Path]:
    """Formats only LibreOffice writes: PPT, XLS, ODS, ODT, RTF."""
    if backends.missing(("soffice",)):
        pytest.skip("needs soffice")
    folder = tmp_path_factory.mktemp("legacy")
    made = {}
    for key, source, fmt in (
        ("ppt", samples["pptx"], "ppt"),
        ("xls", samples["xlsx"], "xls"),
        ("ods", samples["xlsx"], "ods"),
        ("odt", samples["docx"], "odt"),
        ("rtf", samples["docx"], "rtf"),
    ):
        made[key] = _soffice_fixture(source, folder, fmt)
    return made


def _soffice_fixture(source: Path, folder: Path, fmt: str) -> Path:
    exe = backends.require("soffice")
    out = folder / fmt
    out.mkdir(exist_ok=True)
    subprocess.run(
        [
            str(exe),
            f"-env:UserInstallation={(docs._soffice_profile() / 'profile').as_uri()}",
            "--headless",
            "--norestore",
            "--convert-to",
            fmt,
            "--outdir",
            str(out),
            str(source),
        ],
        capture_output=True,
        text=True,
        timeout=600,
        check=True,
    )
    produced = [p for p in out.iterdir() if p.suffix.lstrip(".") == fmt]
    assert produced, f"LibreOffice made no .{fmt} fixture"
    return produced[0]


def _write(path: Path, text: str, encoding: str = "utf-8") -> Path:
    path.write_text(text, encoding=encoding)
    return path


def _make_docx(path: Path) -> Path:
    from docx import Document

    document = Document()
    document.add_heading("Quarterly report", level=1)
    document.add_paragraph("The opening paragraph mentions tangerines.")
    document.add_paragraph("A second paragraph follows it.")
    table = document.add_table(rows=2, cols=2)
    for cell, value in zip(
        (table.cell(0, 0), table.cell(0, 1), table.cell(1, 0), table.cell(1, 1)),
        ("name", "qty", "apple", "3"),
    ):
        cell.text = value
    document.save(path)
    return path


def _make_docx_with_picture(path: Path) -> Path:
    from docx import Document
    from docx.shared import Inches
    from PIL import Image

    picture = path.parent / "swatch.png"
    Image.new("RGB", (120, 80), (200, 60, 60)).save(picture)
    document = Document()
    document.add_heading("Illustrated", level=1)
    document.add_paragraph("Text before the picture.")
    document.add_picture(str(picture), width=Inches(1.5))
    document.save(path)
    return path


def _make_pptx(path: Path, slides: int) -> Path:
    from pptx import Presentation

    deck = Presentation()
    for number in range(1, slides + 1):
        slide = deck.slides.add_slide(deck.slide_layouts[5])
        slide.shapes.title.text = f"Slide {number}"
    deck.save(path)
    return path


def _make_xlsx(path: Path) -> Path:
    from openpyxl import Workbook

    book = Workbook()
    first = book.active
    first.title = "Alpha"
    first.append(["name", "qty"])
    first.append(["apple", 3])
    first.append(["pear", 5])
    second = book.create_sheet("Beta")
    second.append(["x", "y"])
    second.append([1, 2])
    book.save(path)
    return path


def _make_pdf(path: Path, pages: int = 2) -> Path:
    """A PDF with real text and a ruled table, so find_tables has something."""
    import pymupdf

    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 100), "Quarterly report", fontsize=20)
    page.insert_text((72, 140), "The opening line mentions tangerines.", fontsize=11)
    top = 200
    for row in (("name", "qty"), ("apple", "3"), ("pear", "5")):
        left = 72
        for cell in row:
            page.draw_rect(pymupdf.Rect(left, top, left + 100, top + 20))
            page.insert_text((left + 4, top + 14), cell, fontsize=10)
            left += 100
        top += 20
    for number in range(2, pages + 1):
        extra = document.new_page()
        extra.insert_text((72, 100), f"Page {number} has only prose.", fontsize=14)
    document.save(path)
    document.close()
    return path


@pytest.fixture
def work(tmp_path, samples, legacy_or_none):
    """Copy the sources a test needs into its own folder.

    Every runner claims names on disk next to the source, so a test that shared
    the module level fixture folder would see another test's " 2" suffixes.
    """

    def _copy(*keys: str) -> Path | tuple[Path, ...]:
        copied = []
        for key in keys:
            source = samples.get(key) or (legacy_or_none or {})[key]
            target = tmp_path / source.name
            target.write_bytes(source.read_bytes())
            copied.append(target)
        return copied[0] if len(copied) == 1 else tuple(copied)

    return _copy


@pytest.fixture
def legacy_or_none(request):
    """``legacy`` when LibreOffice is there, None otherwise, so a test that
    only wants a PDF does not skip for want of LibreOffice."""
    if backends.missing(("soffice",)):
        return None
    return request.getfixturevalue("legacy")


def run(registry, action: str, source: Path, **params):
    result = registry.run(action, [source], params or None)
    assert result.ok, f"{action} failed: {result.message} ({result.code})"
    assert result.outputs, f"{action} reported no outputs"
    for path in result.outputs:
        assert path.exists(), f"{action} reported a path it did not write: {path}"
        if path.is_file():
            assert path.stat().st_size > 0, f"{action} wrote an empty {path.name}"
    return result


# --- the registry covers the spec table -------------------------------------

#: group, label, the inputs the spec's Conversions table promises. This is the
#: outcome check: a row of the table that lost its entry fails here.
SPEC_ROWS = [
    ("Documents", "DOCX", {"md", "txt", "epub", "rtf", "odt"}),
    ("Documents", "MD", {"docx", "txt", "epub", "rtf", "odt"}),
    ("Documents", "TXT", {"docx", "md", "epub", "rtf", "odt"}),
    ("Documents", "EPUB", {"docx", "md", "txt", "rtf", "odt"}),
    ("Documents", "RTF", {"docx", "md", "txt", "epub", "odt"}),
    ("Documents", "ODT", {"docx", "md", "txt", "epub", "rtf"}),
    ("Documents", "PDF", {"docx", "md", "txt", "epub", "rtf", "odt"}),
    ("Documents", "HTML", {"docx", "md", "txt", "epub", "rtf", "odt"}),
    ("PDF", "DOCX", {"pdf"}),
    ("PDF", "PPTX", {"pdf"}),
    ("PDF", "XLSX", {"pdf"}),
    ("PDF", "JPG", {"pdf"}),
    ("PDF", "PNG", {"pdf"}),
    ("PDF", "EPUB", {"pdf"}),
    ("PDF", "TXT", {"pdf"}),
    ("PDF", "MD", {"pdf"}),
    ("Slides", "PDF", {"pptx", "ppt"}),
    ("Slides", "JPG", {"pptx", "ppt"}),
    ("Slides", "PNG", {"pptx", "ppt"}),
    ("Spreadsheets", "XLSX", {"ods", "csv", "tsv", "xls"}),
    ("Spreadsheets", "ODS", {"xlsx", "csv", "tsv", "xls"}),
    ("Spreadsheets", "CSV", {"xlsx", "ods", "tsv", "xls"}),
    ("Spreadsheets", "TSV", {"xlsx", "ods", "csv", "xls"}),
    ("Spreadsheets", "PDF", {"xlsx", "ods", "csv", "tsv", "xls"}),
    ("Subtitles", "SRT", {"vtt"}),
    ("Subtitles", "VTT", {"srt"}),
    ("Subtitles", "TXT", {"srt", "vtt"}),
]


@pytest.mark.parametrize("group,label,inputs", SPEC_ROWS, ids=lambda v: str(v))
def test_every_spec_conversion_is_a_registry_entry(registry, group, label, inputs):
    matches = [
        a
        for a in registry.all()
        if a.family == CONVERT and a.group == group and a.label == label
    ]
    assert len(matches) == 1, f"{group}/{label}: {[a.id for a in matches]}"
    assert set(matches[0].inputs) == inputs


@pytest.mark.parametrize(
    "label,group,inputs",
    [
        ("Extract Text", "Documents", {"docx", "md", "epub", "rtf", "odt", "html"}),
        ("Extract Markdown", "Documents", {"docx", "txt", "epub", "rtf", "odt", "html"}),
        ("Extract Subtitle Text", "Subtitles", {"srt", "vtt"}),
    ],
)
def test_the_spec_extract_tools_are_registry_entries(registry, label, group, inputs):
    matches = [a for a in registry.all() if a.family == TOOL and a.label == label]
    assert len(matches) == 1
    assert matches[0].group == group
    assert set(matches[0].inputs) == inputs


def test_the_family_loads_clean_and_owns_the_documents_md_segment(registry):
    """text.py handed TXT conversions over; nothing may claim MD twice."""
    assert registry.warnings == []
    assert "text.to_md" not in registry
    assert "txt" in registry.get("docs.doc_to_md").inputs


def test_every_runner_named_by_an_entry_exists(registry):
    for action in registry.all():
        if action.id.startswith("docs."):
            assert callable(action.resolve_runner()), action.id


def test_listing_options_for_a_document_does_not_import_a_backend():
    """``peel options`` must stay instant: the family module may not pull in
    PyMuPDF, python-pptx, openpyxl or python-docx just to list the wheels."""
    code = (
        "import sys, json;"
        "sys.path.insert(0, '.');"
        "from peel_engine.registry import load_registry;"
        "load_registry(force=True).options(['a.pdf', 'b.pdf']);"
        "print(json.dumps(sorted(set(sys.modules) & "
        "{'fitz','pymupdf','pdf2docx','openpyxl','docx','pptx','PIL','numpy'})))"
    )
    proc = subprocess.run(
        [sys.executable, "-c", code], cwd=ENGINE_DIR, capture_output=True, text=True
    )
    assert proc.returncode == 0, proc.stderr
    assert json.loads(proc.stdout) == []


def test_options_on_a_pdf_offers_exactly_the_spec_targets(registry):
    payload = registry.options([Path("whatever.pdf")], include_unavailable=True)
    # Keyed by wheel AND group: the tool wheel has a PDF group of its own
    # (Split, Organize, Compress from the PDF tools bead), and keying by group
    # alone let it overwrite the conversions.
    groups = {
        (wheel["family"], group["group"]): {e["label"] for e in group["entries"]}
        for wheel in payload["wheels"]
        for group in wheel["groups"]
    }
    assert groups[(CONVERT, "PDF")] == {"DOCX", "PPTX", "XLSX", "JPG", "PNG", "EPUB", "TXT", "MD"}


# --- documents --------------------------------------------------------------


@needs_pandoc
@pytest.mark.parametrize("target", ["md", "txt", "html", "epub", "odt"])
def test_docx_converts_to_every_pandoc_document_target(registry, work, target):
    source = work("docx")
    result = run(registry, f"docs.doc_to_{target}", source)
    out = result.outputs[0]
    assert out == source.with_suffix(f".{target}")
    if target in ("md", "txt", "html"):
        assert "tangerines" in out.read_text(encoding="utf-8")
    else:
        assert zipfile.is_zipfile(out)


@needs_office
def test_docx_to_pdf_is_a_pdf_a_reader_opens(registry, work):
    import pymupdf

    result = run(registry, "docs.doc_to_pdf", work("docx"))
    out = result.outputs[0]
    assert out.read_bytes()[:5] == b"%PDF-"
    with pymupdf.open(out) as document:
        assert document.page_count >= 1
        assert "Quarterly report" in document[0].get_text()


@needs_office
def test_docx_to_rtf_is_written_by_libreoffice(registry, work):
    out = run(registry, "docs.doc_to_rtf", work("docx")).outputs[0]
    assert out.read_bytes()[:5] == b"{\\rtf"
    assert "Quarterly report" in out.read_text(encoding="latin-1")


@needs_pandoc
def test_markdown_to_docx_keeps_the_heading(registry, work):
    from docx import Document

    out = run(registry, "docs.doc_to_docx", work("md")).outputs[0]
    document = Document(out)
    assert "Notes" in "\n".join(p.text for p in document.paragraphs)


@needs_office
@pytest.mark.parametrize("target", ["pdf", "rtf"])
def test_a_source_libreoffice_cannot_read_still_reaches_pdf_and_rtf(registry, work, target):
    """Markdown and EPUB take the pandoc-to-ODT bridge. Without it these two
    corners of the Documents row would be missing."""
    out = run(registry, f"docs.doc_to_{target}", work("md")).outputs[0]
    assert out.stat().st_size > 400
    if target == "pdf":
        import pymupdf

        with pymupdf.open(out) as document:
            assert "Notes" in document[0].get_text()


@needs_pandoc
def test_epub_round_trips_back_to_a_document(registry, work):
    epub = run(registry, "docs.doc_to_epub", work("docx")).outputs[0]
    out = run(registry, "docs.doc_to_txt", epub).outputs[0]
    assert "tangerines" in out.read_text(encoding="utf-8")


@needs_pandoc
def test_odt_and_rtf_are_accepted_as_sources(registry, work):
    for key in ("odt", "rtf"):
        out = run(registry, "docs.doc_to_txt", work(key)).outputs[0]
        assert "Quarterly report" in out.read_text(encoding="utf-8")


@needs_pandoc
def test_txt_to_md_keeps_the_words(registry, work):
    out = run(registry, "docs.doc_to_md", work("txt")).outputs[0]
    body = out.read_text(encoding="utf-8")
    assert "First line." in body and "Second line." in body


@needs_pandoc
def test_a_document_with_a_picture_gets_a_markdown_sidecar_folder(registry, work):
    source = work("picdocx")
    result = run(registry, "docs.doc_to_md", source)
    markdown, media = result.outputs
    assert media.is_dir() and media.name == "illustrated-media"
    images = [p for p in media.rglob("*") if p.suffix.lower() == ".png"]
    assert images, "the picture was dropped"
    body = markdown.read_text(encoding="utf-8")
    assert media.name in body
    # the link must be relative, so moving the pair keeps it working
    assert str(source.parent) not in body


@needs_pandoc
def test_a_document_with_no_picture_leaves_no_empty_sidecar_folder(registry, work):
    source = work("docx")
    result = run(registry, "docs.doc_to_md", source)
    assert result.outputs == [source.with_suffix(".md")]
    assert not list(source.parent.glob("*-media"))


@needs_pandoc
def test_html_output_is_one_self_contained_file(registry, work):
    out = run(registry, "docs.doc_to_html", work("picdocx")).outputs[0]
    body = out.read_text(encoding="utf-8")
    assert "<html" in body.lower()
    assert "data:image" in body, "the picture should be embedded, not a sidecar"


@needs_pandoc
@pytest.mark.parametrize(
    "action,tag,ext",
    [
        ("docs.extract_text", "Extract Text", "txt"),
        ("docs.extract_markdown", "Extract Markdown", "md"),
    ],
)
def test_the_extract_tools_name_their_output_after_the_tool(registry, work, action, tag, ext):
    source = work("docx")
    out = run(registry, action, source).outputs[0]
    assert out.name == f"report ({tag}).{ext}"
    assert "tangerines" in out.read_text(encoding="utf-8")


# --- PDF --------------------------------------------------------------------


@needs_pdf2docx
def test_pdf_to_docx_rebuilds_editable_text(registry, work):
    from docx import Document

    out = run(registry, "docs.pdf_to_docx", work("pdf")).outputs[0]
    document = Document(out)
    text = "\n".join(p.text for p in document.paragraphs)
    text += "".join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert "Quarterly report" in text


@needs_pymupdf
def test_pdf_to_pptx_is_one_slide_per_page_at_the_page_shape(registry, work):
    pytest.importorskip("pptx")
    from pptx import Presentation

    out = run(registry, "docs.pdf_to_pptx", work("pdf")).outputs[0]
    deck = Presentation(out)
    assert len(deck.slides) == 2
    assert all(len(slide.shapes) == 1 for slide in deck.slides)
    # A4 portrait in points, within rounding
    assert round(deck.slide_width / 12700) == 595
    assert round(deck.slide_height / 12700) == 842
    picture = deck.slides[0].shapes[0]
    assert picture.width == deck.slide_width


@needs_pymupdf
def test_pdf_to_xlsx_puts_each_page_on_its_own_sheet(registry, work):
    from openpyxl import load_workbook

    result = run(registry, "docs.pdf_to_xlsx", work("pdf"))
    assert "table" in result.message
    book = load_workbook(result.outputs[0])
    assert book.sheetnames == ["Page 1", "Page 2"]
    rows = [[c.value for c in row] for row in book["Page 1"].iter_rows(max_row=3, max_col=2)]
    assert rows == [["name", "qty"], ["apple", 3], ["pear", 5]]
    # a page with no table still carries its text rather than coming out blank
    assert "prose" in str(book["Page 2"]["A1"].value)


@needs_pymupdf
def test_pdf_to_xlsx_writes_numbers_as_numbers_and_words_as_words(registry, work):
    from openpyxl import load_workbook

    book = load_workbook(run(registry, "docs.pdf_to_xlsx", work("pdf")).outputs[0])
    sheet = book["Page 1"]
    assert isinstance(sheet["B2"].value, int)
    assert isinstance(sheet["A2"].value, str)


@needs_pymupdf
@pytest.mark.parametrize("target,magic", [("png", b"\x89PNG"), ("jpg", b"\xff\xd8\xff")])
def test_pdf_to_image_writes_one_numbered_file_per_page(registry, work, target, magic):
    from PIL import Image

    source = work("pdf")
    result = run(registry, f"docs.pdf_to_{target}", source)
    assert [p.name for p in result.outputs] == [
        f"paper (page 1).{target}",
        f"paper (page 2).{target}",
    ]
    for path in result.outputs:
        assert path.read_bytes()[: len(magic)] == magic
        with Image.open(path) as image:
            # A4 at the spec's 300 DPI
            assert image.size == (2480, 3509)


@needs_pymupdf
def test_a_one_page_pdf_converts_to_a_plainly_named_image(registry, work):
    source = work("onepage")
    result = run(registry, "docs.pdf_to_png", source)
    assert [p.name for p in result.outputs] == ["single.png"]


@needs_pymupdf
def test_the_dpi_param_changes_the_raster_size(registry, work):
    from PIL import Image

    out = run(registry, "docs.pdf_to_png", work("onepage"), dpi=72).outputs[0]
    with Image.open(out) as image:
        assert image.size == (595, 842)


@needs_pymupdf
def test_pdf_to_txt_carries_every_page(registry, work):
    body = run(registry, "docs.pdf_to_txt", work("pdf")).outputs[0].read_text(encoding="utf-8")
    assert "Quarterly report" in body
    assert "Page 2 has only prose." in body


@needs_pymupdf
@needs_pandoc
def test_pdf_to_markdown_is_markdown_not_wrapped_html(registry, work):
    body = run(registry, "docs.pdf_to_md", work("pdf")).outputs[0].read_text(encoding="utf-8")
    assert "Quarterly report" in body
    assert "<div" not in body, "the PyMuPDF page wrappers leaked into the Markdown"


@needs_pymupdf
@needs_pandoc
def test_pdf_to_epub_is_an_epub_a_reader_accepts(registry, work):
    out = run(registry, "docs.pdf_to_epub", work("pdf")).outputs[0]
    with zipfile.ZipFile(out) as archive:
        assert archive.read("mimetype") == b"application/epub+zip"


@needs_pymupdf
def test_an_encrypted_pdf_is_a_message_and_leaves_nothing_behind(registry, tmp_path):
    import pymupdf

    locked = tmp_path / "locked.pdf"
    document = pymupdf.open()
    document.new_page().insert_text((72, 100), "secret", fontsize=12)
    document.save(
        locked,
        encryption=pymupdf.PDF_ENCRYPT_AES_256,
        user_pw="letmein",
        owner_pw="letmein",
    )
    document.close()

    result = registry.run("docs.pdf_to_txt", [locked])
    assert result.ok is False
    assert result.code == "pdf_encrypted"
    assert "password" in result.message
    assert sorted(p.name for p in tmp_path.iterdir()) == ["locked.pdf"]


@needs_pymupdf
def test_a_file_that_is_not_really_a_pdf_is_a_message_not_a_traceback(registry, tmp_path):
    fake = tmp_path / "broken.pdf"
    fake.write_bytes(b"this is not a PDF at all")
    result = registry.run("docs.pdf_to_txt", [fake])
    assert result.ok is False
    assert result.code == "bad_pdf"
    assert "broken.pdf" in result.message
    assert sorted(p.name for p in tmp_path.iterdir()) == ["broken.pdf"]


# --- slides -----------------------------------------------------------------


@needs_soffice
def test_pptx_to_pdf_keeps_every_slide(registry, work):
    import pymupdf

    out = run(registry, "docs.slides_to_pdf", work("pptx")).outputs[0]
    with pymupdf.open(out) as document:
        assert document.page_count == 2
        assert "Slide 1" in document[0].get_text()


@needs_soffice
@needs_pymupdf
@pytest.mark.parametrize("target", ["png", "jpg"])
def test_pptx_to_image_writes_one_image_per_slide(registry, work, target):
    from PIL import Image

    result = run(registry, f"docs.slides_to_{target}", work("pptx"))
    assert [p.name for p in result.outputs] == [
        f"deck (slide 1).{target}",
        f"deck (slide 2).{target}",
    ]
    with Image.open(result.outputs[0]) as image:
        assert image.width > 500


@needs_soffice
def test_a_legacy_ppt_is_accepted_as_a_source(registry, work):
    import pymupdf

    out = run(registry, "docs.slides_to_pdf", work("ppt")).outputs[0]
    with pymupdf.open(out) as document:
        assert document.page_count == 2


# --- spreadsheets -----------------------------------------------------------


@needs_soffice
def test_xlsx_to_csv_exports_every_sheet(registry, work):
    """LibreOffice's CSV filter only writes the first sheet, which is why this
    route reads the workbook with openpyxl instead."""
    source = work("xlsx")
    result = run(registry, "docs.sheet_to_csv", source)
    assert [p.name for p in result.outputs] == ["book (Alpha).csv", "book (Beta).csv"]
    first = result.outputs[0].read_bytes()
    assert first.startswith(b"\xef\xbb\xbf"), "Excel reads a BOM-less UTF-8 CSV as cp1252"
    assert first.decode("utf-8-sig").splitlines() == ["name,qty", "apple,3", "pear,5"]


@needs_soffice
def test_xlsx_to_tsv_uses_tabs(registry, work):
    out = run(registry, "docs.sheet_to_tsv", work("xlsx")).outputs[0]
    assert out.read_text(encoding="utf-8-sig").splitlines()[0] == "name\tqty"


def test_csv_and_tsv_convert_to_each_other_without_any_backend(registry, work):
    out = run(registry, "docs.sheet_to_tsv", work("csv")).outputs[0]
    assert out.read_text(encoding="utf-8-sig").splitlines()[1] == "apple\t3\t00731"
    back = run(registry, "docs.sheet_to_csv", work("tsv")).outputs[0]
    assert back.read_text(encoding="utf-8-sig").splitlines()[1] == "apple,3"


def test_csv_to_xlsx_types_the_cells_and_keeps_leading_zeros(registry, work):
    from openpyxl import load_workbook

    result = run(registry, "docs.sheet_to_xlsx", work("csv"))
    sheet = load_workbook(result.outputs[0]).active
    assert sheet["A2"].value == "apple"
    assert sheet["B2"].value == 3
    # 00731 is an identifier, and turning it into 731 is data loss
    assert sheet["C2"].value == "00731"


def test_a_cell_that_starts_with_an_equals_sign_stays_text(registry, tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "odd.csv"
    source.write_text("label,note\n=SUM(A1),ok\n", encoding="utf-8")
    out = run(registry, "docs.sheet_to_xlsx", source).outputs[0]
    cell = load_workbook(out).active["A2"]
    assert cell.value == "=SUM(A1)"
    assert cell.data_type == "s"


def test_a_windows_encoded_csv_is_decoded_not_mangled(registry, tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "accents.csv"
    source.write_bytes("name,city\nJosé,Querétaro\n".encode("cp1252"))
    out = run(registry, "docs.sheet_to_xlsx", source).outputs[0]
    assert load_workbook(out).active["A2"].value == "José"


@needs_soffice
@pytest.mark.parametrize("target", ["ods", "pdf"])
def test_xlsx_converts_to_the_libreoffice_targets(registry, work, target):
    out = run(registry, f"docs.sheet_to_{target}", work("xlsx")).outputs[0]
    if target == "ods":
        assert zipfile.is_zipfile(out)
    else:
        assert out.read_bytes()[:5] == b"%PDF-"


@needs_soffice
def test_an_ods_source_round_trips_to_xlsx(registry, work):
    from openpyxl import load_workbook

    out = run(registry, "docs.sheet_to_xlsx", work("ods")).outputs[0]
    sheet = load_workbook(out)["Alpha"]
    assert [c.value for c in sheet[1]] == ["name", "qty"]


@needs_soffice
def test_a_legacy_xls_source_exports_every_sheet_to_csv(registry, work):
    result = run(registry, "docs.sheet_to_csv", work("xls"))
    assert len(result.outputs) == 2
    assert "apple,3" in result.outputs[0].read_text(encoding="utf-8-sig")


@needs_soffice
def test_csv_to_pdf_goes_through_libreoffice(registry, work):
    out = run(registry, "docs.sheet_to_pdf", work("csv")).outputs[0]
    assert out.read_bytes()[:5] == b"%PDF-"


# --- spreadsheet and document edge cases found by the cross-vendor review ----
#
# Every one of these reproduced before the fix, and every one is silent data
# loss rather than a crash, which is why the happy path tests missed them all.


def test_a_quoted_field_that_spans_two_lines_stays_one_value(registry, tmp_path):
    source = tmp_path / "multi.csv"
    source.write_bytes(b'id,note\r\n1,"first\r\nsecond"\r\n')
    out = run(registry, "docs.sheet_to_tsv", source).outputs[0]
    body = out.read_text(encoding="utf-8-sig")
    assert '"first\nsecond"' in body, "the newline inside the quoted field was eaten"
    assert body.count("\t") == 2, "the field was split across two rows"


def test_a_quoted_field_round_trips_back_through_csv(registry, tmp_path):
    import csv
    import io

    source = tmp_path / "multi.tsv"
    source.write_text('id\tnote\n1\t"first\nsecond"\n', encoding="utf-8")
    out = run(registry, "docs.sheet_to_csv", source).outputs[0]
    rows = list(csv.reader(io.StringIO(out.read_text(encoding="utf-8-sig"), newline="")))
    assert rows == [["id", "note"], ["1", "first\nsecond"]]


def test_brackets_in_a_file_name_do_not_break_the_sheet_name(registry, tmp_path):
    """Windows allows [ ] in a file name; Excel forbids them in a sheet name."""
    from openpyxl import load_workbook

    source = tmp_path / "sales[2026].csv"
    source.write_text("a,b\n1,2\n", encoding="utf-8")
    out = run(registry, "docs.sheet_to_xlsx", source).outputs[0]
    book = load_workbook(out)
    assert book.sheetnames == ["sales 2026"]
    assert book.active["A2"].value == 1


@pytest.mark.parametrize(
    "stem,title",
    [
        ("plan*2026?", "plan 2026"),
        ("a/b:c\\d", "a b c d"),
        ("", "Sheet1"),
        ("[]", "Sheet1"),
        ("x" * 50, "x" * 31),
    ],
)
def test_a_sheet_name_is_always_one_excel_accepts(stem, title):
    assert docs._sheet_title(stem) == title


def test_a_formula_with_no_cached_result_is_exported_not_dropped(registry, tmp_path, monkeypatch):
    """A workbook written by a script and never opened in Excel has formulas
    with no cached value. Exporting a blank cell and reporting success would be
    silent data loss. Without LibreOffice to recalculate, the formula text is
    what is written."""
    from openpyxl import Workbook

    _without_soffice(monkeypatch)
    source = tmp_path / "formula.xlsx"
    book = Workbook()
    book.active["A1"] = "=1+1"
    book.active["B1"] = "label"
    book.save(source)
    out = run(registry, "docs.sheet_to_csv", source).outputs[0]
    assert out.read_text(encoding="utf-8-sig").strip() == "=1+1,label"


def test_a_cached_formula_result_is_preferred_over_the_formula(registry, tmp_path):
    """When Excel did leave a cached value, that value is what belongs in the
    CSV, not the formula text."""
    from openpyxl import Workbook, load_workbook

    import shutil

    source = tmp_path / "cached.xlsx"
    Workbook().save(source)
    # openpyxl cannot write a cached result, so put one in the sheet XML the
    # way Excel does: the formula in <f> and its last computed value in <v>.
    staged = tmp_path / "staged.xlsx"
    cell = (
        b'<sheetData><row r="1"><c r="A1"><f>1+1</f><v>2</v></c></row></sheetData>'
    )
    with zipfile.ZipFile(source) as original, zipfile.ZipFile(staged, "w") as patched:
        for item in original.infolist():
            data = original.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b"<sheetData></sheetData>", cell)
                assert cell in data, "the openpyxl sheet XML changed shape"
            patched.writestr(item, data)
    shutil.move(staged, source)
    assert load_workbook(source, data_only=True).active["A1"].value == 2

    out = run(registry, "docs.sheet_to_csv", source).outputs[0]
    assert out.read_text(encoding="utf-8-sig").strip() == "2"


def test_an_identifier_too_long_for_a_double_stays_text(registry, tmp_path):
    """A spreadsheet holds numbers as doubles, so a 19 digit order number has
    to stay text or it comes back as 1234567890123460000."""
    from openpyxl import load_workbook

    source = tmp_path / "ids.csv"
    source.write_text("id\n1234567890123456789\n", encoding="utf-8")
    out = run(registry, "docs.sheet_to_xlsx", source).outputs[0]
    assert load_workbook(out).active["A2"].value == "1234567890123456789"


@pytest.mark.parametrize(
    "text,expected",
    [
        ("3", 3),
        ("-2.5", -2.5),
        ("00731", "00731"),
        ("1_000", "1_000"),
        ("3kg", "3kg"),
        ("1234567890123456789", "1234567890123456789"),
        ("123456789012345", 123456789012345),
        ("0.5", 0.5),
        ("", ""),
        ("nan", "nan"),
        ("inf", "inf"),
    ],
)
def test_only_a_safe_number_is_written_as_a_number(text, expected):
    assert docs._as_number(text) == expected


@needs_pandoc
def test_an_image_the_source_links_to_relatively_is_not_dropped(registry, tmp_path):
    """pandoc searches its working directory for a relative link, and peel runs
    from wherever the shell started it, not from the document's folder."""
    from PIL import Image

    folder = tmp_path / "elsewhere"
    folder.mkdir()
    Image.new("RGB", (60, 40), (10, 180, 10)).save(folder / "picture.png")
    (folder / "notes.md").write_text("# Hi\n\n![](picture.png)\n", encoding="utf-8")
    out = run(registry, "docs.doc_to_docx", folder / "notes.md").outputs[0]
    with zipfile.ZipFile(out) as archive:
        embedded = [n for n in archive.namelist() if "/media/" in n]
    assert embedded, "the picture beside the source was silently dropped"


# --- subtitles --------------------------------------------------------------


def test_srt_to_vtt_writes_a_header_and_dot_decimals(registry, work):
    out = run(registry, "docs.sub_to_vtt", work("srt")).outputs[0]
    body = out.read_text(encoding="utf-8")
    assert body.startswith("WEBVTT\n\n")
    assert "00:00:01.000 --> 00:00:03.500" in body
    assert "\n1\n" not in body, "SRT index numbers do not belong in a VTT"


def test_vtt_to_srt_numbers_the_cues_and_uses_comma_decimals(registry, work):
    out = run(registry, "docs.sub_to_srt", work("vtt")).outputs[0]
    body = out.read_text(encoding="utf-8")
    assert body.startswith("1\n00:00:01,000 --> 00:00:03,500\nHello there.\n")
    assert "2\n00:00:04,000 --> 00:00:06,250\n" in body
    assert "WEBVTT" not in body
    assert "a comment" not in body, "the NOTE block was read as a cue"
    assert "line:90%" not in body, "the cue settings leaked into the timing line"


def test_a_vtt_short_timecode_is_understood(registry, work):
    """VTT allows MM:SS.mmm with no hour field; SRT never does."""
    cues = docs.read_cues(work("vtt"))
    assert cues[0].start == pytest.approx(1.0)
    assert cues[0].end == pytest.approx(3.5)


def test_subtitles_to_text_drop_the_markup_and_the_repeats(registry, work):
    """The source has three cues: the third repeats the second word for word,
    the way rolling captions do, and a transcript should say it once."""
    out = run(registry, "docs.sub_to_txt", work("srt")).outputs[0]
    body = out.read_text(encoding="utf-8")
    assert body == "Hello there.\nSecond line\nwrapped\n"
    assert "<i>" not in body


def test_the_subtitle_tool_names_its_output_after_the_tool(registry, work):
    out = run(registry, "docs.extract_subtitle_text", work("srt")).outputs[0]
    assert out.name == "movie (Extract Subtitle Text).txt"


def test_a_subtitle_file_with_no_cues_is_a_message_and_writes_nothing(registry, tmp_path):
    empty = tmp_path / "nothing.srt"
    empty.write_text("WEBVTT\n\nNOTE just a note\n", encoding="utf-8")
    result = registry.run("docs.sub_to_vtt", [empty])
    assert result.ok is False
    assert result.code == "empty_input"
    assert sorted(p.name for p in tmp_path.iterdir()) == ["nothing.srt"]


@pytest.mark.parametrize(
    "text,seconds",
    [
        ("00:00:01,500", 1.5),
        ("00:01.250", 1.25),
        ("1:02:03.004", 3723.004),
        ("00:00:02", 2.0),
    ],
)
def test_a_timecode_is_read_in_either_dialect(text, seconds):
    assert docs.parse_timecode(text) == pytest.approx(seconds)


@pytest.mark.parametrize("text", ["", "hello", "-->", "00:00:01,000 --> 00:00:02,000"])
def test_something_that_is_not_a_timecode_is_rejected(text):
    assert docs.parse_timecode(text) is None


def test_a_timecode_is_written_back_the_way_each_format_wants_it():
    assert docs.format_timecode(3723.004, ",") == "01:02:03,004"
    assert docs.format_timecode(0, ".") == "00:00:00.000"


# --- the shared output contract ---------------------------------------------


def test_outputs_land_next_to_the_source(registry, work):
    source = work("srt")
    out = run(registry, "docs.sub_to_txt", source).outputs[0]
    assert out.parent == source.parent


def test_running_twice_never_overwrites(registry, work):
    source = work("srt")
    first = run(registry, "docs.sub_to_txt", source).outputs[0]
    second = run(registry, "docs.sub_to_txt", source).outputs[0]
    assert first.name == "movie.txt"
    assert second.name == "movie 2.txt"
    assert first.read_text(encoding="utf-8") == second.read_text(encoding="utf-8")


@needs_soffice
def test_a_libreoffice_run_leaves_no_scratch_folder_beside_the_file(registry, work):
    source = work("xlsx")
    run(registry, "docs.sheet_to_pdf", source)
    leftovers = [p.name for p in source.parent.iterdir() if p.name.startswith(".peel-")]
    assert leftovers == []


@needs_soffice
def test_the_libreoffice_profile_is_not_written_beside_the_file(registry, work):
    source = work("xlsx")
    run(registry, "docs.sheet_to_ods", source)
    assert docs._soffice_profile().parent != source.parent
    assert sorted(p.name for p in source.parent.iterdir()) == ["book.ods", "book.xlsx"]


def test_a_batch_converts_every_file_and_names_them_apart(registry, work, tmp_path):
    first = work("srt")
    second = tmp_path / "other.srt"
    second.write_bytes(first.read_bytes())
    result = registry.run("docs.sub_to_txt", [first, second])
    assert result.ok
    assert sorted(p.name for p in result.outputs) == ["movie.txt", "other.txt"]


def test_a_batch_where_one_file_is_broken_reports_the_partial(registry, work, tmp_path):
    good = work("srt")
    bad = tmp_path / "broken.srt"
    bad.write_text("not a subtitle at all\n", encoding="utf-8")
    result = registry.run("docs.sub_to_txt", [good, bad])
    assert result.ok is False
    assert result.code == "partial"
    assert [p.name for p in result.outputs] == ["movie.txt"]
    assert "1 of 2" in result.message
    assert not (tmp_path / "broken.txt").exists()


@pytest.mark.parametrize(
    "action,key,wanted",
    [
        ("docs.pdf_to_txt", "docx", "does not accept docx"),
        ("docs.sub_to_vtt", "txt", "does not accept txt"),
        ("docs.sheet_to_pdf", "srt", "does not accept srt"),
        ("docs.slides_to_pdf", "pdf", "does not accept pdf"),
    ],
)
def test_the_wrong_kind_of_file_is_refused_with_a_reason(registry, work, action, key, wanted):
    from peel_engine.result import UnsupportedInput

    with pytest.raises(UnsupportedInput, match=re.escape(wanted)):
        registry.run(action, [work(key)])


def test_an_out_dir_is_honoured(registry, work, tmp_path):
    source = work("srt")
    elsewhere = tmp_path / "chosen"
    result = registry.run("docs.sub_to_txt", [source], out_dir=elsewhere)
    assert result.ok
    assert result.outputs[0] == elsewhere / "movie.txt"


@needs_soffice
def test_a_file_with_no_extension_is_refused_rather_than_guessed(registry, work, tmp_path):
    """The shell hands over whatever Explorer dropped. ``peel options`` offers
    nothing for a bare name, and a direct run says why instead of crashing."""
    from peel_engine.result import UnsupportedInput

    bare = tmp_path / "noextension"
    bare.write_bytes(work("xlsx").read_bytes())
    with pytest.raises(UnsupportedInput, match="no extension"):
        registry.run("docs.sheet_to_pdf", [bare])


def test_a_utf16_subtitle_file_is_decoded(registry, tmp_path):
    """Notepad and several subtitle tools write UTF-16 with a byte order mark."""
    source = tmp_path / "wide.srt"
    source.write_bytes(
        "1\n00:00:01,000 --> 00:00:02,000\nJos\u00e9 habla.\n".encode("utf-16")
    )
    out = run(registry, "docs.sub_to_txt", source).outputs[0]
    assert out.read_text(encoding="utf-8") == "Jos\u00e9 habla.\n"


# --- round 1 adversarial review fixes ----------------------------------------


def _without_soffice(monkeypatch) -> None:
    """Make the engine believe LibreOffice is not installed."""
    real = docs.backends.missing
    monkeypatch.setattr(
        docs.backends, "missing",
        lambda names: [n for n in names if n == "soffice"]
        + list(real(tuple(n for n in names if n != "soffice"))),
    )


@needs_soffice
def test_a_formula_with_no_cached_result_is_recalculated_by_libreoffice(registry, tmp_path):
    from openpyxl import Workbook

    source = tmp_path / "formula.xlsx"
    book = Workbook()
    book.active["A1"] = "=1+1"
    book.active["B1"] = "label"
    book.save(source)
    out = run(registry, "docs.sheet_to_csv", source).outputs[0]
    assert out.read_text(encoding="utf-8-sig").strip() == "2,label"


@pytest.mark.parametrize("office", [True, False])
def test_a_trailing_total_row_with_no_cached_value_is_kept(registry, tmp_path, monkeypatch, office):
    """The trailing all-formula row used to look blank to the value pass and
    was trimmed before the formula pass could see it."""
    from openpyxl import Workbook

    if office and backends.missing(("soffice",)):
        pytest.skip("needs soffice")
    if not office:
        _without_soffice(monkeypatch)
    book = Workbook()
    for row in (1, 2, 3):
        book.active.cell(row, 1, row)
    book.active.cell(4, 1, "=SUM(A1:A3)")
    source = tmp_path / "totals.xlsx"
    book.save(source)
    lines = run(registry, "docs.sheet_to_csv", source).outputs[0].read_text("utf-8-sig").splitlines()
    assert lines == ["1", "2", "3", "6" if office else "=SUM(A1:A3)"]


def test_excel_booleans_and_midnight_dates_are_written_the_way_excel_writes_them(registry, tmp_path):
    import datetime as dt

    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet.append([True, False, dt.datetime(2026, 1, 5), dt.datetime(2026, 1, 5, 13, 30)])
    source = tmp_path / "kinds.xlsx"
    book.save(source)
    out = run(registry, "docs.sheet_to_csv", source).outputs[0]
    assert out.read_text("utf-8-sig").strip() == "TRUE,FALSE,2026-01-05,2026-01-05 13:30:00"


def test_a_semicolon_csv_is_split_into_columns(registry, tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "semi.csv"
    source.write_text("name;qty;price\napple;3;1,50\npear;5;2,25\n", encoding="utf-8")
    rows = list(load_workbook(run(registry, "docs.sheet_to_xlsx", source).outputs[0]).active.values)
    assert rows[0] == ("name", "qty", "price")
    assert rows[1] == ("apple", 3, "1,50")


def test_a_one_column_csv_stays_one_column(registry, tmp_path):
    source = tmp_path / "list.csv"
    source.write_text("name\napple\npear\n", encoding="utf-8")
    assert docs._read_delimited(source, ",") == [["name"], ["apple"], ["pear"]]


def test_srt_cues_without_a_blank_line_between_them_are_still_two_cues(tmp_path):
    source = tmp_path / "tight.srt"
    source.write_text(
        "1\n00:00:01,000 --> 00:00:02,000\nHello\n"
        "2\n00:00:03,000 --> 00:00:04,000\nWorld\n",
        encoding="utf-8",
    )
    cues = docs.read_cues(source)
    assert [(c.start, c.text) for c in cues] == [(1.0, "Hello"), (3.0, "World")]


def test_vtt_cues_without_a_blank_line_between_them_are_still_two_cues(registry, tmp_path):
    source = tmp_path / "tight.vtt"
    source.write_text(
        "WEBVTT\n\n00:00.000 --> 00:01.000\nfirst\n00:01.000 --> 00:02.000\nsecond\n",
        encoding="utf-8",
    )
    assert [c.text for c in docs.read_cues(source)] == ["first", "second"]
    srt = run(registry, "docs.sub_to_srt", source).outputs[0].read_text(encoding="utf-8")
    assert srt.count("-->") == 2 and "00:00:01.000" not in srt


def test_a_number_inside_a_cue_is_kept_when_no_timing_follows(tmp_path):
    source = tmp_path / "count.srt"
    source.write_text("1\n00:00:01,000 --> 00:00:02,000\nCountdown\n3\n\n", encoding="utf-8")
    assert [c.text for c in docs.read_cues(source)] == ["Countdown\n3"]


def test_vtt_voice_and_colour_tags_do_not_reach_an_srt(registry, tmp_path):
    source = tmp_path / "styled.vtt"
    source.write_text(
        "WEBVTT\n\n00:00.000 --> 00:01.000\n<v Roger>Hi <c.yellow>there</c> <i>you</i></v>\n",
        encoding="utf-8",
    )
    srt = run(registry, "docs.sub_to_srt", source).outputs[0].read_text(encoding="utf-8")
    assert "Hi there <i>you</i>" in srt
    assert "<v" not in srt and "<c" not in srt


def _secret(tmp_path: Path) -> tuple[Path, Path]:
    """A private file outside the folder of a downloaded document."""
    private = tmp_path / "private"
    private.mkdir()
    secret = private / "secret.png"
    secret.write_bytes(b"\x89PNG\r\n\x1a\nPRIVATE-KEY-MATERIAL-0123456789")
    downloads = tmp_path / "downloads"
    downloads.mkdir()
    return secret, downloads


def _leaks(path: Path) -> bool:
    import base64

    raw = path.read_bytes()
    needle = b"PRIVATE-KEY-MATERIAL"
    if needle in raw or base64.b64encode(b"\x89PNG\r\n\x1a\nPRIVATE-KEY-MATERIAL")[:24] in raw:
        return True
    if zipfile.is_zipfile(path):
        with zipfile.ZipFile(path) as archive:
            return any(needle in archive.read(n) for n in archive.namelist())
    return False


@needs_pandoc
@pytest.mark.parametrize("action", ["docs.doc_to_html", "docs.doc_to_docx", "docs.doc_to_epub",
                                    "docs.doc_to_odt"])
@pytest.mark.parametrize("link", ["absolute", "parent", "escaped", "raw_html", "css_meta", "file_url"])
def test_a_markdown_file_cannot_pull_in_a_file_outside_its_folder(registry, tmp_path, action, link):
    secret, downloads = _secret(tmp_path)
    reference = {
        "absolute": f"![x]({secret.as_posix()})",
        "parent": "![x](../private/secret.png)",
        "escaped": "![x](..%2Fprivate%2Fsecret.png)",
        "raw_html": f'<img src="{secret.as_posix()}"> <img src=../private/secret.png>',
        "css_meta": "",
        "file_url": f"![x]({secret.as_uri()})",
    }[link]
    front = f"---\ncss: {secret.as_posix()}\n---\n\n" if link == "css_meta" else ""
    source = downloads / "readme.md"
    source.write_text(f"{front}# hi\n\n{reference}\n", encoding="utf-8")
    result = run(registry, action, source)
    assert not _leaks(result.outputs[0]), "a file outside the source folder was embedded"
    assert "outside" in result.message or "Ignored" in result.message, result.message


@needs_office
def test_a_markdown_file_cannot_pull_a_file_into_a_pdf(registry, tmp_path):
    from PIL import Image

    secret, downloads = _secret(tmp_path)
    Image.new("RGB", (40, 40), (200, 0, 0)).save(secret, format="PNG")
    source = downloads / "readme.md"
    source.write_text(f"# hi\n\n![x]({secret.as_posix()})\n", encoding="utf-8")
    result = run(registry, "docs.doc_to_pdf", source)
    import pymupdf

    with pymupdf.open(result.outputs[0]) as pdf:
        assert not pdf[0].get_images(), "an outside image was drawn into the PDF"
    assert "outside" in result.message


@needs_office
def test_a_docx_linking_a_picture_outside_its_folder_does_not_embed_it(registry, tmp_path):
    """A DOCX can link a picture instead of embedding it; LibreOffice would load
    the link when it writes the PDF. Such a document goes through pandoc."""
    import pymupdf
    from PIL import Image

    secret, downloads = _secret(tmp_path)
    Image.new("RGB", (40, 40), (200, 0, 0)).save(secret, format="PNG")
    clean = _make_docx_with_picture(downloads / "base.docx")
    source = downloads / "linked.docx"
    with zipfile.ZipFile(clean) as src, zipfile.ZipFile(source, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "word/_rels/document.xml.rels":
                text = data.decode("utf-8")
                text = re.sub(r'Target="media/[^"]+"',
                              f'Target="{secret.as_uri()}" TargetMode="External"', text)
                data = text.encode("utf-8")
            if item.filename == "word/document.xml":
                data = data.replace(b"r:embed=", b"r:link=")
            out.writestr(item, data)
    assert docs._links_outside_files(source, "docx")
    assert not docs._links_outside_files(clean, "docx")
    result = run(registry, "docs.doc_to_pdf", source)
    with pymupdf.open(result.outputs[0]) as pdf:
        assert not any(page.get_images() for page in pdf)


@needs_pandoc
def test_a_picture_in_a_subfolder_of_the_source_is_kept(registry, tmp_path):
    from PIL import Image

    folder = tmp_path / "doc"
    (folder / "img").mkdir(parents=True)
    Image.new("RGB", (60, 40), (10, 180, 10)).save(folder / "img" / "picture.png")
    (folder / "notes.md").write_text("# Hi\n\n![](img/./picture.png)\n", encoding="utf-8")
    result = run(registry, "docs.doc_to_html", folder / "notes.md")
    assert "data:image/png;base64" in result.outputs[0].read_text(encoding="utf-8")
    assert result.message == f"Wrote {result.outputs[0].name}"


@needs_pandoc
def test_a_web_image_is_not_downloaded_and_stays_a_link(registry, tmp_path):
    import http.server
    import socketserver
    import threading

    hits: list[str] = []

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            hits.append(self.path)
            self.send_response(404)
            self.end_headers()

        def log_message(self, *args):
            pass

    server = socketserver.TCPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_address[1]}/beacon.png"
    try:
        source = tmp_path / "beacon.md"
        source.write_text(f"![logo]({url})\n\n<img src=\"{url}\">\n", encoding="utf-8")
        html = run(registry, "docs.doc_to_html", source)
        docx = run(registry, "docs.doc_to_docx", source)
    finally:
        server.shutdown()
        server.server_close()
    assert hits == []
    assert url in html.outputs[0].read_text(encoding="utf-8")
    assert "web image" in docx.message


def test_document_conversions_offer_the_allow_remote_switch(registry):
    params = {p.name: p for p in registry.get("docs.doc_to_docx").params}
    assert params["allow_remote"].type == "bool" and params["allow_remote"].default is False


def test_a_table_column_empty_in_every_row_is_dropped():
    rows = [["name", "", "qty"], ["apple", None, "3"], ["pear", " ", "5"]]
    assert docs._drop_empty_columns(rows) == [["name", "qty"], ["apple", "3"], ["pear", "5"]]
    assert docs._drop_empty_columns([["a", "", "b"], ["", "x", ""]]) == [["a", "", "b"], ["", "x", ""]]
    assert docs._drop_empty_columns([]) == []


@needs_pymupdf
@pytest.mark.parametrize("action", ["docs.pdf_to_txt", "docs.pdf_to_xlsx"])
def test_a_pdf_with_no_text_says_it_needs_ocr(registry, tmp_path, action):
    pytest.importorskip("openpyxl")
    import pymupdf
    from PIL import Image

    picture = tmp_path / "scan.png"
    Image.new("RGB", (200, 260), (240, 240, 240)).save(picture)
    source = tmp_path / "scan.pdf"
    with pymupdf.open() as pdf:
        page = pdf.new_page()
        page.insert_image(page.rect, filename=str(picture))
        pdf.save(source)
    result = registry.run(action, [source], None)
    assert result.ok and "OCR" in result.message


@needs_pymupdf
def test_a_pdf_with_text_does_not_warn(registry, work):
    assert "OCR" not in run(registry, "docs.pdf_to_txt", work("pdf")).message


@needs_soffice
def test_a_libreoffice_timeout_stops_the_whole_process_tree(registry, tmp_path, monkeypatch):
    """soffice is a launcher for soffice.bin. Killing only the launcher left
    the converter running, and it later wrote a scratch folder beside the
    source after Peel had reported the failure."""
    import time

    monkeypatch.setattr(docs, "SOFFICE_TIMEOUT", 1)
    monkeypatch.setattr(docs, "_SOFFICE_PROFILE", None)  # a slow first run
    source = tmp_path / "a.txt"
    source.write_text("hello\n", encoding="utf-8")
    result = registry.run("docs.doc_to_pdf", [source], None)
    assert not result.ok and result.code == "backend_timeout"
    time.sleep(8)
    assert [p.name for p in tmp_path.iterdir()] == ["a.txt"]


# --- round 3: pandoc runs sandboxed on a staged copy -------------------------


def _png(path: Path, colour=(10, 180, 10)) -> Path:
    from PIL import Image

    path.parent.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", (30, 20), colour).save(path, format="PNG")
    return path


def _epub_media(path: Path) -> list[str]:
    with zipfile.ZipFile(path) as archive:
        return [n for n in archive.namelist() if "/media/" in n]


@needs_pandoc
def test_images_beside_the_source_still_embed_in_html_and_epub(registry, tmp_path):
    folder = tmp_path / "Imágenes de Ñandú"
    _png(folder / "pic.png")
    _png(folder / "raw.png", (200, 10, 10))
    source = folder / "notes.md"
    source.write_text("# Hi\n\n![a](pic.png)\n\n<img src=\"raw.png\">\n", encoding="utf-8")
    html = run(registry, "docs.doc_to_html", source).outputs[0].read_text(encoding="utf-8")
    assert html.count("data:image/png;base64") == 2
    epub = run(registry, "docs.doc_to_epub", source).outputs[0]
    assert len(_epub_media(epub)) == 2


@needs_pandoc
def test_a_stylesheet_inside_the_folder_still_brings_its_own_images(registry, tmp_path):
    folder = tmp_path / "site"
    _png(folder / "css" / "bg.png")
    (folder / "css" / "look.css").write_text("body{background:url(bg.png)}", encoding="utf-8")
    source = folder / "page.md"
    source.write_text('<link rel="stylesheet" href="css/look.css">\n\nhi\n', encoding="utf-8")
    html = run(registry, "docs.doc_to_html", source).outputs[0].read_text(encoding="utf-8")
    assert "data:image/png;base64" in html


@needs_pandoc
@pytest.mark.skipif(os.name == "nt", reason="symlink creation needs privilege on Windows")
def test_a_symlink_inside_the_folder_that_points_out_is_not_followed(registry, tmp_path):
    secret, downloads = _secret(tmp_path)
    os.symlink(secret.parent, downloads / "link")
    source = downloads / "readme.md"
    source.write_text("# hi\n\n![x](link/secret.png)\n", encoding="utf-8")
    for action in ("docs.doc_to_docx", "docs.doc_to_html", "docs.doc_to_epub"):
        assert not _leaks(run(registry, action, source).outputs[0]), action


@needs_pandoc
def test_a_network_path_counts_as_local_even_with_web_images_on(registry, tmp_path):
    """//host/share is a Windows network path. On Linux //tmp/x is /tmp/x, so
    the same reference doubles as a check that the path is never read."""
    secret, downloads = _secret(tmp_path)
    source = downloads / "readme.md"
    double = "/" + secret.as_posix()
    source.write_text(
        f"# hi\n\n![x]({double})\n\n<img src=\"{double}\">\n\n"
        f"![y](\\\\\\\\host\\\\share\\\\x.png)\n",
        encoding="utf-8",
    )
    for action in ("docs.doc_to_html", "docs.doc_to_docx"):
        result = run(registry, action, source, allow_remote=True)
        assert not _leaks(result.outputs[0]), action
        assert "outside" in result.message, result.message


@needs_pandoc
def test_web_images_are_embedded_when_the_run_allows_it(registry, tmp_path):
    import http.server
    import socketserver
    import threading

    picture = _png(tmp_path / "served.png").read_bytes()

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            self.send_response(200)
            self.send_header("Content-Type", "image/png")
            self.send_header("Content-Length", str(len(picture)))
            self.end_headers()
            self.wfile.write(picture)

        def log_message(self, *args):
            pass

    server = socketserver.TCPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    url = f"http://127.0.0.1:{server.server_address[1]}/logo.png"
    try:
        folder = tmp_path / "doc"
        folder.mkdir()
        source = folder / "web.md"
        source.write_text(f"![logo]({url})\n", encoding="utf-8")
        html = run(registry, "docs.doc_to_html", source, allow_remote=True)
    finally:
        server.shutdown()
        server.server_close()
    assert "data:image/png;base64" in html.outputs[0].read_text(encoding="utf-8")


def test_allow_remote_is_offered_only_where_a_web_image_can_land(registry):
    for target in ("docx", "epub", "odt", "rtf", "pdf", "html"):
        names = [p.name for p in registry.get(f"docs.doc_to_{target}").params]
        assert names == ["allow_remote"], target
    for target in ("txt", "md"):
        assert registry.get(f"docs.doc_to_{target}").params == (), target


def test_a_formula_with_a_cached_empty_result_is_a_blank_cell(registry, tmp_path):
    """openpyxl reads a cached "" and a missing cached value both as None. The
    cell type (t="str") tells them apart; no LibreOffice needed."""
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    sheet["A1"], sheet["B1"] = 1, '=IF(A1>5,"big","")'
    raw = tmp_path / "raw.xlsx"
    book.save(raw)
    source = tmp_path / "book.xlsx"
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(source, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                data = data.replace(b'<c r="B1">', b'<c r="B1" t="str">')
                assert b't="str"' in data
            out.writestr(item, data)
    rows = docs._read_workbook(source)
    assert rows[0][1] == [[1, ""]]
    # The same formula with no cached value at all is still not dropped.
    assert docs._read_workbook(raw)[0][1] == [[1, '=IF(A1>5,"big","")']]


def test_a_numeric_last_line_in_a_tight_vtt_cue_is_text(tmp_path):
    source = tmp_path / "room.vtt"
    source.write_text(
        "WEBVTT\n\n00:00.000 --> 00:01.000\nRoom\n101\n00:01.000 --> 00:02.000\nnext\n",
        encoding="utf-8",
    )
    assert [c.text for c in docs.read_cues(source)] == ["Room\n101", "next"]
