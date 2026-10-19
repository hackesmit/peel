"""Verifier failing checks for hq-plct.4, narrow final fix 22c90b1.

Not collected by default (the name does not start with test_). Run:
    cd engine && python3 -m pytest -q tests/verify_hq_plct_4.py

B1 (blocking): docs._xml_kind looks only at the first 256 bytes for "<". A
relationships part that starts with more than 256 bytes of whitespace and no
XML declaration is legal XML, LibreOffice reads it, but the gate classifies it
as a binary part and skips it. The external picture link in it survives, and
the private picture from outside the folder is embedded in the output.

N1 (non-blocking note, silent incompleteness): when the gate removes a link
LibreOffice would have loaded, the result message is still only "Wrote x.pdf".
The pandoc routes name what they left out; the LibreOffice routes do not.
"""
from __future__ import annotations

import re
from pathlib import Path

import pytest

from test_docs import (  # noqa: F401  (registry is a fixture)
    registry, _red_secret, _swatch, _rezip, needs_soffice, needs_pymupdf,
)


def _red_in_pdf(pdf: Path) -> int:
    import pymupdf

    total = 0
    with pymupdf.open(pdf) as doc:
        for page in doc:
            pix = page.get_pixmap(dpi=40)
            raw, n = pix.samples, pix.n
            total += sum(1 for i in range(0, len(raw), n)
                         if raw[i] > 150 and raw[i + 1] < 70 and raw[i + 2] < 70)
    return total


def _padded(text: str) -> bytes:
    """The same relationships part, declaration dropped, 400 newlines first."""
    return b"\n" * 400 + re.sub(r"^<\?xml[^>]*\?>", "", text).encode("utf-8")


def _xlsx(downloads: Path, secret: Path, pad: bool) -> Path:
    from openpyxl import Workbook
    from openpyxl.drawing.image import Image as SheetImage

    book = Workbook()
    book.active["A1"] = "total"
    book.active.add_image(SheetImage(str(_swatch(downloads))), "B2")
    clean = downloads / "base.xlsx"
    book.save(clean)

    def edit(name: str, data: bytes) -> bytes:
        if name == "xl/drawings/_rels/drawing1.xml.rels":
            text = re.sub(r'Target="[^"]*media/[^"]+"',
                          f'Target="{secret.as_uri()}" TargetMode="External"', data.decode("utf-8"))
            return _padded(text) if pad else text.encode("utf-8")
        if name == "xl/drawings/drawing1.xml":
            return data.replace(b"r:embed=", b"r:link=")
        return data

    return _rezip(clean, downloads / "linked.xlsx", edit)


def _pptx(downloads: Path, secret: Path) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_picture(str(_swatch(downloads)), Inches(1), Inches(1), Inches(6), Inches(4))
    clean = downloads / "base.pptx"
    deck.save(clean)

    def edit(name: str, data: bytes) -> bytes:
        if name == "ppt/slides/_rels/slide1.xml.rels":
            text = re.sub(r'Target="\.\./media/[^"]+"',
                          f'Target="{secret.as_uri()}" TargetMode="External"', data.decode("utf-8"))
            return _padded(text)
        if name == "ppt/slides/slide1.xml":
            return data.replace(b"r:embed=", b"r:link=")
        return data

    return _rezip(clean, downloads / "linked.pptx", edit)


@needs_soffice
@needs_pymupdf
def test_b1_xlsx_rels_after_leading_whitespace_cannot_embed_an_outside_picture(registry, tmp_path):
    pytest.importorskip("openpyxl")
    secret, downloads = _red_secret(tmp_path)
    source = _xlsx(downloads, secret, pad=True)
    result = registry.run("docs.sheet_to_pdf", [source], None)
    if result.ok:
        assert _red_in_pdf(result.outputs[0]) == 0, "the outside picture is in the PDF"


@needs_soffice
@needs_pymupdf
def test_b1_pptx_rels_after_leading_whitespace_cannot_embed_an_outside_picture(registry, tmp_path):
    pytest.importorskip("pptx")
    secret, downloads = _red_secret(tmp_path)
    source = _pptx(downloads, secret)
    result = registry.run("docs.slides_to_pdf", [source], None)
    if result.ok:
        assert _red_in_pdf(result.outputs[0]) == 0, "the outside picture is in the PDF"


@needs_soffice
@needs_pymupdf
def test_n1_a_removed_link_is_named_in_the_result(registry, tmp_path):
    """Non-blocking note: the gate removes the picture link, the message says nothing."""
    pytest.importorskip("openpyxl")
    secret, downloads = _red_secret(tmp_path)
    source = _xlsx(downloads, secret, pad=False)
    result = registry.run("docs.sheet_to_pdf", [source], None)
    assert result.ok and _red_in_pdf(result.outputs[0]) == 0
    assert result.message.strip() != f"Wrote {result.outputs[0].name}", (
        "a linked picture was removed but the result does not say so")
