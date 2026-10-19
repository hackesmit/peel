"""Adversarial review checks for hq-plct.4 (docs family), round 1.

Collected with the rest of the suite (renamed from review_hq_plct_4.py once fixed):

    cd engine && python3 -m pytest -q tests/test_review_hq_plct_4.py

Each test is a failing check against b2f154b..7bd1b81. BLOCK tests are the
blocking findings; NOTE tests document non-blocking findings and are marked
xfail(strict=False) so they report without failing the run.
"""

from __future__ import annotations

import base64
import os
import subprocess
import sys
import time
from pathlib import Path

import pytest

from peel_engine import backends
from peel_engine.families import docs
from peel_engine.registry import load_registry


def needs(*names: str):
    missing = backends.missing(tuple(names))
    return pytest.mark.skipif(bool(missing), reason=f"needs {', '.join(missing)}")


@pytest.fixture(scope="module")
def registry():
    return load_registry(force=True)


# --- BLOCK 1: trailing formula row with no cached value is dropped ----------


@needs("openpyxl")
def test_block_xlsx_to_csv_keeps_a_trailing_total_row_with_no_cached_value(registry, tmp_path):
    """A script-written workbook (never opened in Excel) whose LAST row is a
    formula. 391ae9e says uncached formulas are recovered, but the data_only
    pass trims the trailing all-blank row first, no '' cell remains, the
    second pass never runs, and the total row silently disappears."""
    from openpyxl import Workbook

    book = Workbook()
    sheet = book.active
    for row in (1, 2, 3):
        sheet.cell(row, 1, row)
    sheet.cell(4, 1, "=SUM(A1:A3)")
    source = tmp_path / "totals.xlsx"
    book.save(source)

    result = registry.run("docs.sheet_to_csv", [source], None)
    assert result.ok
    lines = result.outputs[0].read_text("utf-8-sig").splitlines()
    assert len(lines) == 4, f"total row lost: {lines}"


# --- BLOCK 2: a cue without a blank line swallows the next cue --------------


def test_block_srt_without_blank_line_between_cues_is_two_cues(tmp_path):
    source = tmp_path / "tight.srt"
    source.write_text(
        "1\n00:00:01,000 --> 00:00:02,000\nHello\n"
        "2\n00:00:03,000 --> 00:00:04,000\nWorld\n",
        encoding="utf-8",
    )
    cues = docs.read_cues(source)
    assert [c.text for c in cues] == ["Hello", "World"], cues


def test_block_vtt_without_blank_line_between_cues_is_two_cues(tmp_path):
    source = tmp_path / "tight.vtt"
    source.write_text(
        "WEBVTT\n\n00:00.000 --> 00:01.000\nfirst\n00:01.000 --> 00:02.000\nsecond\n",
        encoding="utf-8",
    )
    cues = docs.read_cues(source)
    assert [c.text for c in cues] == ["first", "second"], cues


def test_block_srt_to_vtt_never_emits_an_srt_timecode_as_cue_text(registry, tmp_path):
    source = tmp_path / "tight.srt"
    source.write_text(
        "1\n00:00:01,000 --> 00:00:02,000\nHello\n"
        "2\n00:00:03,000 --> 00:00:04,000\nWorld\n",
        encoding="utf-8",
    )
    result = registry.run("docs.sub_to_vtt", [source], None)
    assert result.ok
    assert "00:00:03,000" not in result.outputs[0].read_text(encoding="utf-8")


# --- BLOCK 3: local files outside the source folder are packed into output --


@needs("pandoc")
@pytest.mark.parametrize("action", ["docs.doc_to_html", "docs.doc_to_docx"])
def test_block_markdown_cannot_embed_a_file_from_outside_its_folder(registry, tmp_path, action):
    """A downloaded Markdown file names an absolute path (on George:
    C:\\Users\\danie\\...). Converting it must not copy that file's bytes into
    the HTML or DOCX the user will then send to someone."""
    private = tmp_path / "private"
    private.mkdir()
    secret = private / "secret.png"
    secret.write_bytes(b"\x89PNG\r\n\x1a\nPRIVATE-KEY-MATERIAL-0123456789")
    downloads = tmp_path / "downloads"
    downloads.mkdir()
    source = downloads / "readme.md"
    source.write_text(f"# hi\n\n![x]({secret.as_posix()})\n", encoding="utf-8")

    result = registry.run(action, [source], None)
    assert result.ok
    out = result.outputs[0].read_bytes()
    needle_b64 = base64.b64encode(secret.read_bytes())[:24]
    if action.endswith("html"):
        assert needle_b64 not in out, "secret file embedded as a data URI"
    else:
        import zipfile

        with zipfile.ZipFile(result.outputs[0]) as z:
            blobs = b"".join(z.read(n) for n in z.namelist() if "media" in n)
        assert b"PRIVATE-KEY-MATERIAL" not in blobs, "secret file embedded in word/media"


# --- BLOCK 4 (George only): PDF to XLSX on the pymupdf setup.ps1 installs ---
# Not a new test: the DoD's own tests fail on George.
#   rsync worktree to /mnt/george/peel-dev/review-4d15/, then on George with a
#   venv holding pymupdf 2.0 + pymupdf_layout 2.0 (a plain `pip install pymupdf`
#   today, which is what install/setup.ps1 runs, unpinned):
#   python -m pytest -q tests/test_docs.py -k pdf_to_xlsx   -> 2 failed
#   find_tables().extract() gives [['name','','qty'], ...]: a phantom empty
#   column, so every table lands shifted one column right.


@needs("pymupdf", "openpyxl")
def test_block_pdf_to_xlsx_has_no_all_empty_column_inside_a_table(registry, tmp_path):
    """Version independent form of the George failure: no column of a written
    table may be entirely empty. Passes on hqbox (pymupdf 1.26), fails on
    George (pymupdf 2.0)."""
    sys.path.insert(0, str(Path(__file__).parent))
    import test_docs
    from openpyxl import load_workbook

    pdf = test_docs._make_pdf(tmp_path / "paper.pdf")
    result = registry.run("docs.pdf_to_xlsx", [pdf], None)
    assert result.ok
    sheet = load_workbook(result.outputs[0])["Page 1"]
    for column in sheet.iter_cols(min_row=1, max_row=3, values_only=True):
        assert any(v not in (None, "") for v in column), list(sheet.values)


# --- NOTES (non-blocking) ---------------------------------------------------


@needs("pandoc")
@pytest.mark.xfail(strict=False, reason="NOTE: remote fetch is a policy call for Daniel")
def test_note_converting_markdown_does_not_hit_the_network(registry, tmp_path):
    import http.server
    import socketserver
    import threading

    hits: list[str] = []

    class Handler(http.server.BaseHTTPRequestHandler):
        def do_GET(self):  # noqa: N802
            hits.append(self.path)
            self.send_response(404)
            self.end_headers()

        def log_message(self, *args):
            pass

    server = socketserver.TCPServer(("127.0.0.1", 0), Handler)
    threading.Thread(target=server.serve_forever, daemon=True).start()
    try:
        source = tmp_path / "beacon.md"
        source.write_text(
            f"![x](http://127.0.0.1:{server.server_address[1]}/beacon.png)\n", encoding="utf-8"
        )
        registry.run("docs.doc_to_docx", [source], None)
    finally:
        server.shutdown()
    assert hits == [], f"pandoc fetched {hits}"


@needs("soffice")
@pytest.mark.skipif(sys.platform == "win32", reason="Windows kills soffice.bin with the launcher")
@pytest.mark.xfail(strict=False, reason="NOTE: Linux only, George did not reproduce")
def test_note_soffice_timeout_leaves_no_orphan_writing_next_to_the_source(registry, tmp_path, monkeypatch):
    monkeypatch.setattr(docs, "SOFFICE_TIMEOUT", 1)
    monkeypatch.setattr(docs, "_SOFFICE_PROFILE", None)  # force a slow first run
    source = tmp_path / "a.txt"
    source.write_text("hello\n", encoding="utf-8")
    result = registry.run("docs.doc_to_pdf", [source], None)
    assert not result.ok and result.code == "backend_timeout"
    time.sleep(8)
    stray = [p.name for p in tmp_path.iterdir() if p.name.startswith(".peel-")]
    assert stray == [], f"orphaned soffice.bin recreated scratch output: {stray}"


@pytest.mark.xfail(strict=False, reason="NOTE: semicolon CSV (es-ES Excel) is not sniffed")
@needs("openpyxl")
def test_note_semicolon_csv_is_split_into_columns(registry, tmp_path):
    from openpyxl import load_workbook

    source = tmp_path / "semi.csv"
    source.write_text("name;qty;price\napple;3;1,50\n", encoding="utf-8")
    result = registry.run("docs.sheet_to_xlsx", [source], None)
    rows = list(load_workbook(result.outputs[0]).active.values)
    assert rows[0] == ("name", "qty", "price"), rows
